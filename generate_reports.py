import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_table_borders(table, color="B0BEC5"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'  <w:top w:val="single" w:sz="6" w:space="0" w:color="{color}"/>\n'
        f'  <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="{color}"/>\n'
        f'  <w:left w:val="single" w:sz="6" w:space="0" w:color="{color}"/>\n'
        f'  <w:right w:val="single" w:sz="6" w:space="0" w:color="{color}"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def format_paragraph(p, space_before=0, space_after=6, line_spacing=1.15):
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing

def add_title_page(doc, work_type, topic, variant=None):
    # Header
    p = doc.add_paragraph()
    format_paragraph(p, 0, 2, 1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("ТОО «КОЛЛЕДЖ ХЕКСЛЕТ КАЗАХСТАН»\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True

    p = doc.add_paragraph()
    format_paragraph(p, 0, 48, 1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Отделение информационных технологий и программирования\nСпециальность 4S06130103 «Разработчик программного обеспечения»")
    r.font.name = "Times New Roman"
    r.font.size = Pt(11)
    r.font.italic = True

    # Work type
    p = doc.add_paragraph()
    format_paragraph(p, 48, 6, 1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(work_type.upper())
    r.font.name = "Times New Roman"
    r.font.size = Pt(16)
    r.font.bold = True
    r.font.color.rgb = RGBColor(26, 35, 126)

    p = doc.add_paragraph()
    format_paragraph(p, 0, 10, 1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(f"по дисциплине «Применение информационно-коммуникационных и цифровых технологий»\nТема: «{topic}»")
    r.font.name = "Times New Roman"
    r.font.size = Pt(13)
    r.font.bold = True

    if variant:
        p = doc.add_paragraph()
        format_paragraph(p, 0, 16, 1.15)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{variant}")
        r.font.name = "Times New Roman"
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(13, 71, 161)

    p = doc.add_paragraph()
    format_paragraph(p, 0, 96, 1.15)

    # Info block
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(3.2)
    table.columns[1].width = Inches(3.3)

    cell_right = table.cell(0, 1)
    p = cell_right.paragraphs[0]
    format_paragraph(p, 0, 2, 1.15)
    r = p.add_run("Выполнила:\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

    r = p.add_run("студентка группы 21WD\nПилипенко Алина\n\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    r = p.add_run("Проверил:\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)
    r.font.bold = True

    r = p.add_run("преподаватель ИКТ, магистр управления и ИБ\nКалиаскаров Д.А.\n\n")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    r = p.add_run("Оценка: ___________ / 100 б.\nПодпись: ___________")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    # Footer
    p = doc.add_paragraph()
    format_paragraph(p, 100, 0, 1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Алматы, 2026 г.")
    r.font.name = "Times New Roman"
    r.font.size = Pt(12)

    doc.add_page_break()

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, 14, 4, 1.15)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(13, 71, 161)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    format_paragraph(p, 10, 3, 1.15)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(33, 33, 33)
    return p

def add_body(doc, text, bold_prefix=None):
    p = doc.add_paragraph()
    format_paragraph(p, 0, 4, 1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    return p

def add_bullet(doc, text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    format_paragraph(p, 0, 3, 1.15)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = "Times New Roman"
        r_pre.font.size = Pt(11.5)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r.font.size = Pt(11.5)
    return p

def add_code_block(doc, code_text):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.5)
    cell = tbl.cell(0, 0)
    set_cell_background(cell, "F5F5F5")
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    set_table_borders(tbl, color="CFD8DC")
    p = cell.paragraphs[0]
    format_paragraph(p, 0, 0, 1.05)
    r = p.add_run(code_text)
    r.font.name = "Consolas"
    r.font.size = Pt(9.5)
    r.font.color.rgb = RGBColor(38, 50, 56)
    p_after = doc.add_paragraph()
    format_paragraph(p_after, 0, 3, 1.15)

def build_lab_report():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.79)
        s.bottom_margin = Inches(0.79)
        s.left_margin = Inches(0.98)
        s.right_margin = Inches(0.79)

    add_title_page(doc, "ОТЧЕТ ПО ПРАКТИЧЕСКОЙ РАБОТЕ", "Окружение и структура проекта", "Вариант 5: REST API-сервис (код, миграции БД, тесты, документация)")

    add_heading_1(doc, "Введение и цель работы")
    add_body(doc, "Цель работы: освоить принципы взаимодействия операционной системы, аппаратного обеспечения и файловой системы; спроектировать переносимую и безопасную структуру проекта; изучить правила построения абсолютных и относительных путей; освоить методы устранения платформозависимых ошибок (бага «работает только у меня») через контейнеризацию и стандарты именования.")
    add_body(doc, "Объект разработки: серверный REST API-сервис с версионированием базы данных, автоматизированным тестированием и документацией интерфейсов.")

    add_heading_1(doc, "1. Проектирование структуры папок проекта (Вариант 5)")
    add_body(doc, "Для реализации надежного REST API-сервиса спроектирована модульная структура по принципам Clean Architecture и стандартам 12-Factor App. Все имена папок и файлов соответствуют критериям надежности: использована только латиница, нижний регистр (lowercase), дефисы и знаки подчеркивания вместо пробелов.")

    tree_text = """rest-api-service/
├── app/                      # Исходный код сервиса (FastAPI/SQLAlchemy)
│   ├── api/                  # Контроллеры и эндпоинты
│   │   └── v1/               # Версионирование API v1
│   │       ├── auth.py       # Эндпоинты авторизации и регистрации
│   │       └── users.py      # CRUD операции над пользователями
│   ├── core/                 # Ядро приложения
│   │   ├── config.py         # Настройки и валидация .env (Pydantic Settings)
│   │   └── security.py       # Хэширование паролей (bcrypt), токены JWT
│   ├── models/               # Описание таблиц БД (SQLAlchemy ORM)
│   │   └── user.py           # Модель таблицы пользователей
│   ├── schemas/              # Pydantic DTO (валидация входных/выходных данных)
│   │   └── user.py           # Схемы запросов и ответов
│   ├── services/             # Слой бизнес-логики сервиса
│   │   └── user_service.py   # Логика регистрации и аутентификации
│   ├── __init__.py
│   └── main.py               # Точка входа приложения (ASGI instance)
├── data/                     # Данные проекта
│   └── initial_data.csv      # Начальные справочники для наполнения БД (seeding)
├── docs/                     # Документация сервиса
│   ├── api.md                # Описание интеграции и сценариев использования
│   └── openapi.json          # Экспортированная спецификация OpenAPI/Swagger
├── migrations/               # Управление схемой базы данных (Alembic)
│   ├── versions/             # Изолированные скрипты миграций
│   │   └── 001_initial_tables.py
│   ├── alembic.ini           # Конфигурация инструмента Alembic
│   └── env.py                # Скрипт подключения Alembic к моделям app
├── tests/                    # Автоматизированные тесты (pytest)
│   ├── conftest.py           # Фикстуры тестовой сессии и тестовая БД
│   ├── test_api.py           # Интеграционные тесты эндпоинтов
│   └── test_auth.py          # Тестирование безопасности и JWT-токенов
├── .dockerignore             # Исключения контекста сборки Docker
├── .env.example              # Шаблон конфигурации окружения (без секретов)
├── .gitignore                # Исключения Git (секреты, кэш, .venv)
├── docker-compose.yml        # Оркестрация сервиса и БД PostgreSQL
├── Dockerfile                # Спецификация сборки рабочего образа сервиса
├── pyproject.toml            # Декларативное описание зависимостей и инструментов
└── README.md                 # Руководство по развертыванию и запуску"""
    add_code_block(doc, tree_text)

    add_heading_2(doc, "Подробное функциональное назначение элементов:")
    add_bullet(doc, "содержит всю логику серверного приложения. Разделение на core, api, models, schemas и services обеспечивает слабую связанность компонентов (Loose Coupling) и высокую тестируемость.", "app/: ")
    add_bullet(doc, "обеспечивает воспроизводимое версионирование схемы реляционной БД. Каждое изменение таблиц фиксируется в атомарном файле миграции, что исключает конфликт структуры данных при совместной разработке.", "migrations/: ")
    add_bullet(doc, "содержит модульные и интеграционные тесты на базе pytest. Позволяет автоматически проверять работоспособность эндпоинтов и бизнес-логики до деплоя.", "tests/: ")
    add_bullet(doc, "содержит спецификацию OpenAPI (Swagger) и руководство по интеграции, обеспечивая прозрачный контракт взаимодействия с фронтенд-командой.", "docs/: ")
    add_bullet(doc, "хранит CSV-фикстуры для первоначального наполнения таблиц БД начальными справочниками.", "data/: ")
    add_bullet(doc, "гарантируют одинаковое развертывание сервиса и СУБД PostgreSQL на машинах всех участников команды и на сервере.", "Dockerfile и docker-compose.yml: ")
    add_bullet(doc, "предотвращает утечку конфиденциальных данных (.env) и мусорных артефактов в удаленный репозиторий.", ".gitignore: ")

    add_heading_1(doc, "2. Обоснование выбора операционных систем")
    add_body(doc, "Для разработки выбрана операционная система Windows 11 с подсистемой WSL2 (Ubuntu 24.04 LTS), что позволяет разработчику использовать привычный графический интерфейс, современные десктопные IDE (VS Code, JetBrains) и инструменты тестирования, сохраняя при этом прямое взаимодействие с нативным ядром Linux и идентичной файловой системой.", "ОС для разработки: ")
    add_body(doc, "Для боевого сервера (production) выбрана операционная система Linux (Ubuntu Server 24.04 LTS или легковесный Alpine Linux в Docker-контейнере), поскольку она обеспечивает минимальные накладные расходы на системные ресурсы за счет отсутствия графической оболочки (headless), наивысшую стабильность аптайма и строгую модель безопасности ядра.", "ОС для сервера: ")

    add_heading_1(doc, "3. Анализ путей внутри проекта")
    add_body(doc, "В рамках спроектированного дерева каталогов определены 3 типа путей к файлам проекта:")

    add_body(doc, "/home/developer/projects/rest-api-service/app/main.py — полный путь от корня файловой системы (/), однозначно идентифицирующий файл независимо от рабочего каталога. В исходном коде такой путь не используется, так как он жестко привязан к конкретному хосту и пользователю.", "1. Абсолютный путь: ")
    add_body(doc, "./migrations/versions/001_initial_tables.py — путь от корня проекта к файлу начальной миграции. Символ ./ обозначает текущий рабочий каталог; путь корректно работает на любой машине, куда склонирован репозиторий.", "2. Относительный путь №1: ")
    add_body(doc, "../app/core/config.py — путь указан относительно директории tests/. Маркер ../ обеспечивает переход в родительский каталог с последующим обращением к модулю конфигурации ядра, гарантируя переносимость тестов без хардкода абсолютных адресов.", "3. Относительный путь №2: ")

    add_heading_1(doc, "4. Анализ и исправление «плохих» имен файлов")
    add_body(doc, "В исходном перечне обнаружены типичные ошибки именования, нарушающие переносимость и безопасность. Ниже приведено сопоставление ошибок и их исправление под архитектуру Варианта 5:")

    tbl = doc.add_table(rows=4, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(1.5)
    tbl.columns[1].width = Inches(1.5)
    tbl.columns[2].width = Inches(1.5)
    tbl.columns[3].width = Inches(2.0)
    set_table_borders(tbl)

    headers = ["Исходное имя", "Обнаруженные ошибки", "Исправленное имя", "Техническое обоснование"]
    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        set_cell_background(cell, "E3F2FD")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        format_paragraph(p, 0, 0, 1.05)
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)
        r.font.bold = True

    rows_data = [
        ("Мои Файлы/Главный Скрипт.py",
         "1. Кириллица в пути и имени\n2. Пробелы между словами\n3. Заглавные буквы",
         "app/main.py\n(или scripts/seed_database.py)",
         "Пробелы в терминале парсятся как разделители параметров. Кириллица вызывает сбои локали на Linux-серверах. В проекте REST API главный модуль именуется main.py в папке app/."),
        ("C:\\temp\\Data.CSV",
         "1. Жесткий абсолютный путь\n2. Обратные слеши \\\n3. Заглавные буквы в имени\n4. Расширение в ВЕРХНЕМ регистре",
         "data/initial_data.csv",
         "Каталог C:\\temp отсутствует на Linux-серверах и машинах коллег. Расширение .CSV из-за чувствительности Linux к регистру не совпадет с маской *.csv. Заменено на относительный путь data/initial_data.csv."),
        ("тест 1.js",
         "1. Кириллица\n2. Пробел\n3. Расширение .js вместо .py",
         "tests/test_auth.py\n(или tests/test_api.py)",
         "Для сервиса на Python тесты пишутся на Python (.py). Кириллица и пробел ломают вызов pytest. Фреймворк pytest требует имена вида test_*.py для автоматического поиска тестов.")
    ]

    for r_idx, row in enumerate(rows_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx + 1, c_idx)
            set_cell_margins(cell, 60, 60, 80, 80)
            if r_idx % 2 == 1:
                set_cell_background(cell, "FAFAFA")
            p = cell.paragraphs[0]
            format_paragraph(p, 0, 0, 1.05)
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(10)

    add_heading_1(doc, "5. Продвинутая часть: Единое окружение команды")
    add_body(doc, "Для полной синхронизации окружения команды внедрен стек на основе Docker, docker-compose и подсистемы WSL2, дополненный строгой фиксацией зависимостей в pyproject.toml.")
    add_heading_2(doc, "Почему это устраняет баг «работает только у меня»:")
    add_bullet(doc, "Файловая система Windows (NTFS) не различает регистр букв: код с ошибочным регистром data/Users.csv локально сработает на Windows, но завершится фатальным сбоем FileNotFoundError на боевом сервере Linux (ext4). Запуск в Docker-контейнере или WSL2 немедленно выявляет такие баги на этапе локальной разработки.", "1. Различие регистрозависимости файловых систем: ")
    add_bullet(doc, "В Windows стандартным окончанием строк является CRLF (\\r\\n), а в Linux — LF (\\n). Bash-скрипты, сохраненные с CRLF, падают в Linux с ошибкой «^M: bad interpreter». Контейнеризация и настройка .gitattributes исключают эту проблему.", "2. Разница разделителей путей и перевода строк: ")
    add_bullet(doc, "Разные минорные версии Python (3.10 vs 3.12) или отличия в компиляторах C для нативных библиотек базы данных (psycopg2/libpq) приводят к несовместимости байткода. Образ Docker гарантирует одинаковую версию рантайма и системных зависимостей у каждого разработчика.", "3. Разночтение версий библиотек и системных пакетов: ")
    add_bullet(doc, "Команда docker compose up поднимает идентичную инфраструктуру (сервис, СУБД PostgreSQL) с заранее заданными сетевыми портами и переменными окружения, независимо от базовой ОС хоста.", "4. Полная воспроизводимость окружения: ")

    add_heading_1(doc, "6. Вывод")
    add_body(doc, "Спроектированная архитектура проекта REST API полностью удовлетворяет критериям масштабируемости, безопасности и переносимости. Использование валидных правил именования, относительных путей и контейнеризации гарантирует детерминированное поведение приложения в любой инфраструктуре.")

    doc.save("Отчет_Занятие_02_Вариант_5_Пилипенко_Алина.docx")
    print("Lab report docx successfully generated!")

def build_srs_report():
    doc = Document()
    for s in doc.sections:
        s.top_margin = Inches(0.79)
        s.bottom_margin = Inches(0.79)
        s.left_margin = Inches(0.98)
        s.right_margin = Inches(0.79)

    add_title_page(doc, "САМОСТОЯТЕЛЬНАЯ РАБОТА СТУДЕНТА (СРС)", "Сравнительный анализ операционных систем Windows и Linux", None)

    add_heading_1(doc, "1. Введение и актуальность исследования")
    add_body(doc, "Операционная система (ОС) представляет собой базовый комплекс программного обеспечения, выступающий посредником между аппаратным обеспечением ЭВМ, прикладными программами и пользователем. ОС выполняет ключевые функции диспетчеризации процессов, распределения оперативной памяти, организации файловых структур и разграничения прав доступа.")
    add_body(doc, "В современной индустрии информационных технологий доминирующие позиции занимают семейства операционных систем Microsoft Windows и GNU/Linux. Выбор операционной системы является определяющим фактором эффективности рабочего процесса инженера-программиста, устойчивости развертывания веб-сервисов и безопасности критической инфраструктуры. Цель данной работы — комплексный сравнительный анализ Windows и Linux по 7 фундаментальным критериям и формирование аргументированного вывода по выбору ОС.")

    add_heading_1(doc, "2. Сравнительная таблица Windows и Linux")

    tbl = doc.add_table(rows=8, cols=4)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(1.3)
    tbl.columns[1].width = Inches(1.8)
    tbl.columns[2].width = Inches(1.8)
    tbl.columns[3].width = Inches(1.6)
    set_table_borders(tbl)

    headers = ["Критерий", "Microsoft Windows", "GNU/Linux", "Инженерный вывод"]
    for c_idx, h in enumerate(headers):
        cell = tbl.cell(0, c_idx)
        set_cell_background(cell, "E3F2FD")
        set_cell_margins(cell, 80, 80, 100, 100)
        p = cell.paragraphs[0]
        format_paragraph(p, 0, 0, 1.05)
        r = p.add_run(h)
        r.font.name = "Times New Roman"
        r.font.size = Pt(10.5)
        r.font.bold = True

    srs_data = [
        ("1. Модель лицензии и исходного кода",
         "Проприетарное коммерческое ПО с закрытым исходным кодом. Платное лицензирование клиентских и серверных копий.",
         "Свободное и открытое ПО (Open Source, GPL/MIT). Исходный код доступен для аудита, модификации и бесплатного развертывания.",
         "Linux полностью исключает лицензионные затраты при масштабировании инфраструктуры."),
        ("2. Файловая система и регистр",
         "Логические диски (C:\\, D:\\). Обратные слеши (\\). Регистронезависима (case-insensitive). Основная ФС — NTFS.",
         "Единое иерархическое дерево от корня (/). Прямой слеш (/). Строго регистрозависима (case-sensitive). ФС: ext4, Btrfs, XFS.",
         "Необходимость строгого использования нижнего регистра в коде во избежание сбоев при переносе."),
        ("3. Потребление аппаратных ресурсов",
         "Высокое потребление RAM (от 3-4 ГБ для ОС) из-за неудаляемого графического интерфейса и телеметрии.",
         "Высокоэффективный режим headless (без GUI). Сервер запускается при объеме от 256-512 МБ RAM.",
         "Linux отдает максимум мощностей аппаратуры прикладным серверным сервисам."),
        ("4. Информационная безопасность",
         "Широкая поверхность атаки. Массовая мишень для вирусов и шифровальщиков. UAC часто отключается пользователями.",
         "Архитектура Security by Design. Строгая модель прав (chmod/chown, root, sudo). Механизмы SELinux, AppArmor, cgroups.",
         "Linux обеспечивает высокий уровень защиты в сетевых средах и закрытых контурах."),
        ("5. Управление пакетами и средой",
         "Ручная установка бинарных установщиков (.exe, .msi). Проблема конфликтов библиотек (DLL Hell).",
         "Централизованные репозитории и пакетные менеджеры (apt, dnf, pacman, nix). Декларативная установка.",
         "В Linux воспроизводимость окружения автоматизируется одной строкой скрипта."),
        ("6. Командная строка и автоматизация",
         "Акцент на графический интерфейс. PowerShell функционален, но имеет высокий порог входа и медленный старт.",
         "Терминал — основной инструмент. Стандартные утилиты POSIX, мощные конвейеры (pipes |), скрипты Bash/Zsh.",
         "Автоматизация сборки и деплоя на серверах Linux реализуется значительно проще и быстрее."),
        ("7. Контейнеризация и экосистема",
         "Docker работает через виртуализацию (Hyper-V / WSL2), создавая дополнительные накладные расходы.",
         "Нативная среда для Docker и Kubernetes (изоляция на уровне ячеек ядра cgroups/namespaces). Доля на серверах > 85%.",
         "Стандарт де-факто для бэкенд-разработки, облачных микросервисов и DevOps.")
    ]

    for r_idx, row in enumerate(srs_data):
        for c_idx, val in enumerate(row):
            cell = tbl.cell(r_idx + 1, c_idx)
            set_cell_margins(cell, 60, 60, 80, 80)
            if r_idx % 2 == 1:
                set_cell_background(cell, "FAFAFA")
            p = cell.paragraphs[0]
            format_paragraph(p, 0, 0, 1.05)
            r = p.add_run(val)
            r.font.name = "Times New Roman"
            r.font.size = Pt(9.5)

    add_heading_1(doc, "3. Аргументированный вывод о выборе ОС")
    add_body(doc, "Проведенное исследование позволяет сделать вывод, что в современной разработке программного обеспечения противопоставление Windows и Linux уступило место гибридной синергетической модели:")
    add_bullet(doc, "оптимальным выбором является Windows 11 с установленной подсистемой WSL2 (Ubuntu). Это позволяет совместить преимущества привычной настольной ОС с богатым выбором прикладного ПО и комфортом графических IDE с полноценным доступом к нативному терминалу и ядру Linux.", "Для локальной рабочей станции разработчика: ")
    add_bullet(doc, "безальтернативным стандартом выступает Linux (Ubuntu Server / Debian / Alpine). Преимущества в виде минимального потребления памяти, высочайшей отказоустойчивости, встроенных механизмов безопасности (SELinux) и нативной поддержки контейнеров делают Linux единственным выбором для размещения REST API, микросервисов и СУБД.", "Для серверной инфраструктуры (production): ")

    add_heading_1(doc, "4. Примечание об использовании инструментов ИИ")
    add_body(doc, "В соответствии с критериями приёмки СРС, подтверждается использование инструментов искусственного интеллекта для структурирования сравнительных данных. Полученные аналитические выкладки верифицированы автором, сопоставлены с официальной документацией ядра Linux и Microsoft Learn, обезличены и адаптированы под требования дисциплины.")

    add_heading_1(doc, "5. Список использованных источников")
    add_bullet(doc, "Шыныбеков Д.А., Ускенбаева Р.К. Информационно-коммуникационные технологии: Учебник. — Алматы: АУЭС, 2017.")
    add_bullet(doc, "Гаврилов М.В., Климов В.А. Информатика и информационные технологии. — М.: Юрайт, 2020.")
    add_bullet(doc, "Закон Республики Казахстан «Об информатизации» от 24.11.2015 № 418-V.")
    add_bullet(doc, "ГОСО ТиПО. Приказ Министра просвещения РК № 348 от 03.08.2022.")
    add_bullet(doc, "Microsoft Learn: Документация подсистемы WSL2. URL: https://learn.microsoft.com/ru-ru/windows/wsl/")
    add_bullet(doc, "The Linux Kernel documentation. URL: https://www.kernel.org/doc/html/latest/")

    doc.save("СРС_Сравнительный_анализ_Windows_и_Linux_Пилипенко_Алина.docx")
    print("SRS report docx successfully generated!")

if __name__ == "__main__":
    build_lab_report()
    build_srs_report()
